"""Regenerate small, valid Office/PDF visual fixtures (fixture-build deps optional)."""
from pathlib import Path
import pymupdf as fitz
from docx import Document
from docx.shared import Inches as DocInches
from pptx import Presentation
from pptx.util import Inches as PptInches

HERE = Path(__file__).parent
pix = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 20, 20), 0)
pix.clear_with(0x00ff00)
png = HERE / 'green-square.png'
png.write_bytes(pix.tobytes('png'))

word = Document()
heading = word.add_paragraph(style='Heading 1')
heading.add_run('Hello ').bold = True
heading.add_run('world 2026').italic = True
table = word.add_table(rows=1, cols=1)
table.cell(0, 0).text = 'Table cell'
word.add_picture(str(png), width=DocInches(.3))
word.save(HERE / 'sample.docx')

slides = Presentation()
slide = slides.slides.add_slide(slides.slide_layouts[6])
shape = slide.shapes.add_textbox(PptInches(1), PptInches(1), PptInches(3), PptInches(1))
shape.text_frame.text = 'Hello'
slide.shapes.add_picture(str(png), PptInches(6), PptInches(4), width=PptInches(.5))
slides.save(HERE / 'sample.pptx')

pdf = fitz.open()
page = pdf.new_page(width=400, height=300)
page.draw_rect(fitz.Rect(20, 20, 380, 280), color=(1, 0, 0), width=2)
page.insert_text((30, 60), 'Hello world', fontsize=14, color=(0.1, 0.2, 0.7))
page.insert_link({'kind': fitz.LINK_URI, 'from': fitz.Rect(30, 240, 120, 260), 'uri': 'https://example.invalid/'})
page.insert_image(fitz.Rect(350, 250, 365, 265), stream=png.read_bytes())
pdf.save(HERE / 'sample.pdf', garbage=3, deflate=True)
pdf.close()
